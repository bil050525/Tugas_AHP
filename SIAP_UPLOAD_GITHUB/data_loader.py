"""
Modul Pembaca Data Excel & Ekspor/Impor
Sistem Pendukung Keputusan Pemilihan Mobil (AHP)

Modul ini bertanggung jawab untuk:
1. Membaca file Excel studi kasus asli (AHP_Studi_Kasus_Mobil.xlsx)
2. Mengekstrak matriks perbandingan kriteria dan alternatif
3. Memvalidasi struktur data dan format matriks
4. Mendukung upload file Excel baru dan perhitungan ulang dinamis
5. Mengekspor hasil perankingan dan perhitungan ke format CSV dan Excel
"""

import os
import io
import csv
import openpyxl
import pandas as pd
import numpy as np

# Path default data Excel
DEFAULT_EXCEL_PATH = os.path.join(os.path.dirname(__file__), "data", "AHP_Studi_Kasus_Mobil.xlsx")

# Definisi nama kriteria lengkap & singkat
DEFAULT_CRITERIA_MAP = {
    "S": "Style (Gaya/Tampilan)",
    "K": "Keandalan (Reliability/Durabilitas)",
    "B": "Konsumsi Bahan Bakar (Efisiensi BBM)"
}

DEFAULT_CRITERIA_KEYS = ["S", "K", "B"]
DEFAULT_CRITERIA_LABELS = [DEFAULT_CRITERIA_MAP[k] for k in DEFAULT_CRITERIA_KEYS]
DEFAULT_ALTERNATIVES = ["Avanza", "Xenia", "Ertiga", "Grand Livina"]


def get_default_dataset() -> dict:
    """
    Menyediakan data fallback bawaan yang 100% presisi sesuai dokumen soal Excel.
    Digunakan sebagai cadangan jika file Excel tidak dapat diakses.
    """
    criteria_matrix = [
        [1.0, 0.5, 3.0],
        [2.0, 1.0, 4.0],
        [1.0 / 3.0, 0.25, 1.0]
    ]

    alt_matrices = {
        "S": [
            [1.0, 2.0, 1.0 / 3.0, 0.25],
            [0.5, 1.0, 0.25, 0.20],
            [3.0, 4.0, 1.0, 2.0],
            [4.0, 5.0, 0.5, 1.0]
        ],
        "K": [
            [1.0, 2.0, 0.25, 1.0 / 3.0],
            [0.5, 1.0, 0.20, 0.25],
            [4.0, 5.0, 1.0, 3.0],
            [3.0, 4.0, 1.0 / 3.0, 1.0]
        ],
        "B": [
            [1.0, 2.0, 0.5, 4.0],
            [0.5, 1.0, 2.0 / 3.0, 3.0],
            [2.0, 1.5, 1.0, 4.0],
            [0.25, 1.0 / 3.0, 0.25, 1.0]
        ]
    }

    return {
        "source": "Default Data (AHP Studi Kasus Mobil)",
        "criteria_keys": DEFAULT_CRITERIA_KEYS,
        "criteria_labels": DEFAULT_CRITERIA_LABELS,
        "alternatives": DEFAULT_ALTERNATIVES,
        "criteria_matrix": criteria_matrix,
        "alt_matrices": alt_matrices
    }


def load_data_from_excel(filepath: str = DEFAULT_EXCEL_PATH) -> dict:
    """
    Membaca dan memparsing data studi kasus dari file Excel AHP_Studi_Kasus_Mobil.xlsx
    Mengambil matriks perbandingan kriteria dan alternatif secara dinamis.
    """
    if not os.path.exists(filepath):
        # Jika file belum ada, gunakan dataset default yang presisi
        return get_default_dataset()

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet_name = "AHP_Mobil" if "AHP_Mobil" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # 1. Baca Matriks Kriteria (Baris 18-20, Kolom B-D)
        criteria_matrix = []
        for r in range(18, 21):
            row_vals = []
            for c in range(2, 5):
                val = ws.cell(row=r, column=c).value
                row_vals.append(float(val) if val is not None else 1.0)
            criteria_matrix.append(row_vals)

        # 2. Baca Alternatif (Baris 32, Kolom B-E)
        alternatives = []
        for c in range(2, 6):
            alt_name = ws.cell(row=32, column=c).value
            if alt_name:
                alternatives.append(str(alt_name).strip())
            else:
                alternatives.append(DEFAULT_ALTERNATIVES[c - 2])

        # 3. Baca Matriks Style (Baris 33-36, Kolom B-E)
        style_matrix = []
        for r in range(33, 37):
            row_vals = []
            for c in range(2, 6):
                val = ws.cell(row=r, column=c).value
                row_vals.append(float(val) if val is not None else 1.0)
            style_matrix.append(row_vals)

        # 4. Baca Matriks Keandalan (Baris 50-53, Kolom B-E)
        keandalan_matrix = []
        for r in range(50, 54):
            row_vals = []
            for c in range(2, 6):
                val = ws.cell(row=r, column=c).value
                row_vals.append(float(val) if val is not None else 1.0)
            keandalan_matrix.append(row_vals)

        # 5. Baca Matriks Konsumsi BBM (Baris 67-70, Kolom B-E)
        bbm_matrix = []
        for r in range(67, 71):
            row_vals = []
            for c in range(2, 6):
                val = ws.cell(row=r, column=c).value
                row_vals.append(float(val) if val is not None else 1.0)
            bbm_matrix.append(row_vals)

        alt_matrices = {
            "S": style_matrix,
            "K": keandalan_matrix,
            "B": bbm_matrix
        }

        return {
            "source": os.path.basename(filepath),
            "criteria_keys": DEFAULT_CRITERIA_KEYS,
            "criteria_labels": DEFAULT_CRITERIA_LABELS,
            "alternatives": alternatives,
            "criteria_matrix": criteria_matrix,
            "alt_matrices": alt_matrices
        }

    except Exception as e:
        print(f"Peringatan: Gagal membaca file Excel ({e}), menggunakan dataset default.")
        return get_default_dataset()


def validate_matrix(matrix: list, expected_size: int = None) -> tuple:
    """
    Validasi integritas matriks perbandingan berpasangan AHP:
    - Harus berupa matriks bujur sangkar (n x n)
    - Elemen diagonal utama harus bernilai 1.0
    - Elemen a[i][j] > 0
    - Disarankan a[j][i] approx 1 / a[i][j]
    """
    if not isinstance(matrix, list) or len(matrix) == 0:
        return False, "Matriks kosong atau bukan berformat list"

    n = len(matrix)
    if expected_size and n != expected_size:
        return False, f"Ukuran matriks ({n}) tidak sesuai yang diharapkan ({expected_size})"

    for i in range(n):
        if len(matrix[i]) != n:
            return False, f"Baris {i+1} memiliki {len(matrix[i])} kolom, seharusnya {n}"
        for j in range(n):
            val = matrix[i][j]
            try:
                val_f = float(val)
                if val_f <= 0:
                    return False, f"Nilai perbandingan pada baris {i+1}, kolom {j+1} harus lebih dari 0"
                if i == j and abs(val_f - 1.0) > 1e-4:
                    return False, f"Elemen diagonal utama baris {i+1} harus bernilai 1.0 (ditemukan: {val_f})"
            except (ValueError, TypeError):
                return False, f"Nilai pada baris {i+1}, kolom {j+1} bukan angka yang valid"

    return True, "Valid"


def export_results_to_csv(results_data: dict) -> str:
    """
    Mengekspor hasil perankingan dan skor akhir mobil ke format CSV string.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow(["HASIL SISTEM PENDUKUNG KEPUTUSAN PEMILIHAN MOBIL (AHP)"])
    writer.writerow([])
    
    # Kriteria dan Bobot
    writer.writerow(["BOBOT PRIORITAS KRITERIA"])
    writer.writerow(["No", "Kriteria", "Bobot Desimal", "Persentase (%)"])
    for i, (crit_name, weight) in enumerate(results_data["criteria_weights"].items(), 1):
        writer.writerow([i, crit_name, f"{weight:.6f}", f"{weight*100:.2f}%"])
    writer.writerow([])

    # Hasil Perankingan Mobil
    writer.writerow(["PERINGKAT AKHIR MOBIL"])
    header_rank = ["Peringkat", "Nama Mobil", "Skor Akhir", "Persentase (%)"]
    # Tambahkan kolom per kriteria
    for c_name in results_data["criteria_weights"].keys():
        header_rank.append(f"Kontribusi {c_name}")
    writer.writerow(header_rank)

    for item in results_data["ranked_alternatives"]:
        row = [
            item["rank"],
            item["name"],
            f"{item['final_score']:.6f}",
            f"{item['final_score_percent']:.2f}%"
        ]
        for c_name in results_data["criteria_weights"].keys():
            b_val = item["breakdown"].get(c_name, {}).get("weighted_score", 0.0)
            row.append(f"{b_val:.6f}")
        writer.writerow(row)
    writer.writerow([])

    # Kesimpulan Rekomendasi
    best = results_data.get("best_alternative")
    if best:
        writer.writerow(["REKOMENDASI MOBIL TERBAIK"])
        writer.writerow(["Mobil Terpilih", best["name"]])
        writer.writerow(["Skor Tertinggi", f"{best['final_score']:.6f} ({best['final_score_percent']:.2f}%)"])
    writer.writerow([])

    # Ringkasan Uji Konsistensi
    writer.writerow(["RINGKASAN UJI KONSISTENSI (AHP)"])
    writer.writerow(["Matriks", "Lambda Max", "CI", "RI", "CR", "Status (CR <= 0.10)"])
    for c_info in results_data.get("consistency_summary", []):
        writer.writerow([
            c_info["matrix_name"],
            f"{c_info['lambda_max']:.6f}",
            f"{c_info['ci']:.6f}",
            f"{c_info['ri']:.2f}",
            f"{c_info['cr']:.6f}",
            c_info["status"]
        ])

    return output.getvalue()


def export_results_to_excel(results_data: dict) -> io.BytesIO:
    """
    Mengekspor hasil perankingan dan perhitungan ke format workbook Excel (.xlsx)
    yang rapi dan berformat profesional.
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Peringkat & Rekomendasi
    ws_rank = wb.active
    ws_rank.title = "Peringkat & Rekomendasi"

    ws_rank.append(["HASIL SISTEM PENDUKUNG KEPUTUSAN PEMILIHAN MOBIL (AHP)"])
    ws_rank.append([])
    
    best = results_data.get("best_alternative")
    if best:
        ws_rank.append(["REKOMENDASI MOBIL TERBAIK:", best["name"]])
        ws_rank.append(["SKOR AKHIR:", f"{best['final_score']:.6f} ({best['final_score_percent']:.2f}%)"])
        ws_rank.append([])

    # Tabel Peringkat
    header = ["Peringkat", "Nama Mobil", "Skor Akhir", "Persentase"]
    for c_name in results_data["criteria_weights"].keys():
        header.append(f"Bobot {c_name}")
    ws_rank.append(header)

    for item in results_data["ranked_alternatives"]:
        row = [
            item["rank"],
            item["name"],
            round(item["final_score"], 6),
            f"{item['final_score_percent']:.2f}%"
        ]
        for c_name in results_data["criteria_weights"].keys():
            w_val = item["breakdown"].get(c_name, {}).get("weighted_score", 0.0)
            row.append(round(w_val, 6))
        ws_rank.append(row)

    ws_rank.append([])
    # Tabel Bobot Kriteria
    ws_rank.append(["BOBOT KRITERIA PENILAIAN"])
    ws_rank.append(["Kriteria", "Bobot Desimal", "Bobot (%)"])
    for crit_name, w in results_data["criteria_weights"].items():
        ws_rank.append([crit_name, round(w, 6), f"{w*100:.2f}%"])

    # Sheet 2: Uji Konsistensi
    ws_cons = wb.create_sheet(title="Uji Konsistensi")
    ws_cons.append(["RINGKASAN UJI KONSISTENSI MATRIKS PERBANDINGAN BERPASANGAN"])
    ws_cons.append([])
    ws_cons.append(["Matriks", "Lambda Max", "CI", "RI", "CR", "Status", "Keterangan"])

    for c_info in results_data.get("consistency_summary", []):
        ws_cons.append([
            c_info["matrix_name"],
            round(c_info["lambda_max"], 6),
            round(c_info["ci"], 6),
            round(c_info["ri"], 2),
            round(c_info["cr"], 6),
            c_info["status"],
            "CR <= 0.10 (Konsisten)" if c_info["is_consistent"] else "CR > 0.10 (Harus Diperbaiki)"
        ])

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
